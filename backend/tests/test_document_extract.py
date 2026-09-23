import io
import unittest
import zipfile

from docx import Document

from app.document_extract import DocumentExtractionError, extract_document


def simple_pdf(text=b"Hoc lai DEMO-1"):
    stream=b"BT /F1 12 Tf 72 720 Td ("+text+b") Tj ET"
    objects=[
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length "+str(len(stream)).encode()+b" >>\nstream\n"+stream+b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    result=bytearray(b"%PDF-1.4\n")
    offsets=[0]
    for number,obj in enumerate(objects,1):
        offsets.append(len(result));result.extend(f"{number} 0 obj\n".encode()+obj+b"\nendobj\n")
    xref=len(result);result.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:result.extend(f"{offset:010d} 00000 n \n".encode())
    result.extend(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return bytes(result)


class DocumentExtractionTests(unittest.TestCase):
    def test_extract_pdf_text_and_metadata(self):
        result=extract_document("policy.pdf",simple_pdf())
        self.assertEqual(result["format"],"pdf")
        self.assertEqual(result["page_count"],1)
        self.assertIn("Hoc lai DEMO-1",result["content"])
        self.assertEqual(len(result["sha256"]),64)

    def test_extract_docx_paragraphs_and_tables(self):
        document=Document();document.add_heading("Quy trinh DEMO",1);document.add_paragraph("Noi dung hoc lai")
        table=document.add_table(rows=1,cols=2);table.cell(0,0).text="Muc";table.cell(0,1).text="Gia tri"
        buffer=io.BytesIO();document.save(buffer)
        result=extract_document("guide.docx",buffer.getvalue())
        self.assertEqual(result["format"],"docx")
        self.assertIn("Noi dung hoc lai",result["content"])
        self.assertIn("Muc | Gia tri",result["content"])

    def test_rejects_unsupported_and_unsafe_docx(self):
        with self.assertRaisesRegex(DocumentExtractionError,"UNSUPPORTED"):
            extract_document("policy.doc",b"legacy")
        buffer=io.BytesIO()
        with zipfile.ZipFile(buffer,"w") as archive:
            archive.writestr("[Content_Types].xml","x")
            archive.writestr("../escape.xml","x")
        with self.assertRaisesRegex(DocumentExtractionError,"UNSAFE"):
            extract_document("unsafe.docx",buffer.getvalue())

    def test_rejects_empty_text_pdf(self):
        from pypdf import PdfWriter
        writer=PdfWriter();writer.add_blank_page(width=100,height=100)
        buffer=io.BytesIO();writer.write(buffer)
        with self.assertRaisesRegex(DocumentExtractionError,"TEXT_NOT_FOUND"):
            extract_document("scan.pdf",buffer.getvalue())

    def test_chunk_markdown_rich_citation_labels(self):
        from advisor_core.rag import chunk_markdown
        text = "## Trang 1\n\nKhoan 1: Dieu kien tot nghiep.\n\nKhoan 2: Chung chi tieng Anh.\n\n## Trang 2\n\nKhoan 3: Diem ren luyen."
        chunks = chunk_markdown(text, document_id="doc-1", version_id="v-1", scope="academic_policy", source="demo", valid_from="2026-09-01")
        self.assertEqual(len(chunks), 5)
        self.assertEqual(chunks[1]["section"], "Trang 1, Đoạn 2")
        self.assertEqual(chunks[2]["section"], "Trang 1, Đoạn 3")
        self.assertEqual(chunks[4]["section"], "Trang 2, Đoạn 5")

    def test_image_extraction_mock(self):
        from unittest.mock import patch
        with patch("app.document_extract._ocr_image", return_value="Quy che dao tao tu anh chup"):
            fake_png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 100
            res = extract_document("scan_test.png", fake_png)
            self.assertEqual(res["format"], "image")
            self.assertEqual(res["page_count"], 1)
            self.assertIn("Quy che dao tao", res["content"])


if __name__=="__main__":
    unittest.main()

